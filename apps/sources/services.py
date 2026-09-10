import hashlib
import json
import re
from datetime import date, datetime, time
from decimal import Decimal
from zipfile import BadZipFile

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.core.audit import audit_event
from apps.people.models import (
    AdministrativeSystem,
    EmploymentStatus,
    Person,
    PersonSystemMembership,
)
from apps.people.validators import normalize_curp, validate_curp

from .models import ImportStatus, PayrollImport, PayrollSourceRow, SourceRowStatus


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _normalized(value) -> str:
    return str(value or "").strip().casefold()


def _configured_values(mapping, key):
    return {_normalized(value) for value in mapping.get(key, []) if _normalized(value)}


def _detect_systems(payload, mapping):
    unified = set()
    flags = set()
    warnings = []
    federal_values = _configured_values(mapping, "federal_values")
    state_values = _configured_values(mapping, "state_values")

    systems_header = mapping.get("systems")
    if systems_header:
        raw = payload.get(systems_header)
        tokens = {_normalized(raw)}
        tokens.update(_normalized(token) for token in re.split(r"[,;/|+]", str(raw or "")))
        if federal_values.intersection(tokens):
            unified.add("FEDERAL")
        if state_values.intersection(tokens):
            unified.add("ESTATAL")
        if raw not in (None, "") and not unified:
            warnings.append("El valor de sistema no coincide con el mapeo configurado.")

    for code, column_key, values_key in (
        ("FEDERAL", "federal_flag", "federal_values"),
        ("ESTATAL", "state_flag", "state_values"),
    ):
        header = mapping.get(column_key)
        if header and _normalized(payload.get(header)) in _configured_values(mapping, values_key):
            flags.add(code)

    has_unified_value = bool(systems_header and str(payload.get(systems_header, "") or "").strip())
    has_flag_values = any(
        mapping.get(key) and str(payload.get(mapping[key], "") or "").strip()
        for key in ("federal_flag", "state_flag")
    )
    if has_unified_value and has_flag_values and unified != flags:
        warnings.append("Las columnas configuradas de sistema contienen valores contradictorios.")
    return sorted(unified | flags), warnings


def _employment_status(raw, mapping):
    value = _normalized(raw)
    if value and value in _configured_values(mapping, "active_values"):
        return EmploymentStatus.ACTIVE
    if value and value in _configured_values(mapping, "inactive_values"):
        return EmploymentStatus.INACTIVE
    return EmploymentStatus.UNKNOWN


def _hash_uploaded_file(uploaded_file):
    digest = hashlib.sha256()
    uploaded_file.seek(0)
    for chunk in uploaded_file.chunks():
        digest.update(chunk)
    uploaded_file.seek(0)
    return digest.hexdigest()


def _open_workbook(uploaded_file, mapping):
    uploaded_file.seek(0)
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError) as exc:
        raise ValidationError(f"No fue posible leer el archivo .xlsx: {exc}") from exc
    sheet_name = mapping.get("sheet_name")
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValidationError(f"La hoja '{sheet_name}' no existe en el archivo.")
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.active
    return workbook, sheet


@transaction.atomic
def import_payroll_xlsx(*, uploaded_file, mapping, actor):
    if not actor.has_perm("sources.import_payroll"):
        raise PermissionError("No tiene permiso para importar la fuente oficial.")
    if not uploaded_file.name.lower().endswith(".xlsx"):
        raise ValidationError("Solamente se admiten archivos .xlsx.")
    if uploaded_file.size > settings.MAX_IMPORT_FILE_BYTES:
        raise ValidationError("El archivo excede el límite configurado.")

    sha256 = _hash_uploaded_file(uploaded_file)
    if PayrollImport.objects.filter(sha256=sha256).exists():
        raise ValidationError("Este archivo ya fue importado; coincide su hash SHA-256.")

    workbook, sheet = _open_workbook(uploaded_file, mapping)
    header_row = int(mapping.get("header_row", 1))
    header_cells = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in header_cells]
    nonempty_headers = [header for header in headers if header]
    if len(nonempty_headers) != len(set(nonempty_headers)):
        workbook.close()
        raise ValidationError("La fila de encabezados contiene nombres repetidos o vacíos intermedios.")

    required = mapping.get("curp")
    configured_headers = [
        mapping.get(key) for key in ("curp", "name", "employment_status", "systems", "federal_flag", "state_flag")
        if mapping.get(key)
    ]
    missing = [header for header in configured_headers if header not in headers]
    if not required or missing:
        workbook.close()
        detail = ", ".join(missing) if missing else "CURP"
        raise ValidationError(f"No se encontraron los encabezados configurados: {detail}.")

    uploaded_file.seek(0)
    import_record = PayrollImport.objects.create(
        source_file=uploaded_file,
        original_filename=uploaded_file.name[:255],
        sha256=sha256,
        imported_by=actor,
        column_mapping=mapping,
    )

    seen_curps = set()
    seen_rows = set()
    counts = {"rows": 0, "accepted": 0, "errors": 0, "warnings": 0}
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        payload = {header: _json_value(value) for header, value in zip(headers, values) if header}
        if not any(value not in (None, "") for value in payload.values()):
            continue
        counts["rows"] += 1
        errors = []
        warnings = []
        raw_curp = str(payload.get(mapping["curp"]) or "").strip()
        curp = normalize_curp(raw_curp)
        try:
            validate_curp(curp)
        except ValidationError as exc:
            errors.extend(exc.messages)

        fingerprint = hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
        ).hexdigest()
        if fingerprint in seen_rows:
            errors.append("Fila duplicada dentro del mismo archivo.")
        seen_rows.add(fingerprint)
        if curp and curp in seen_curps:
            errors.append("CURP duplicada dentro del mismo archivo.")
        if curp:
            seen_curps.add(curp)

        systems, system_warnings = _detect_systems(payload, mapping)
        warnings.extend(system_warnings)
        full_name = str(payload.get(mapping.get("name"), "") or "").strip() if mapping.get("name") else ""
        status_raw = (
            str(payload.get(mapping.get("employment_status"), "") or "").strip()
            if mapping.get("employment_status")
            else ""
        )
        person = None
        if not errors:
            person, created = Person.objects.get_or_create(
                curp=curp,
                defaults={"full_name": full_name, "employment_status": _employment_status(status_raw, mapping)},
            )
            previous = {
                "full_name": person.full_name,
                "employment_status": person.employment_status,
            }
            changed = False
            if full_name and person.full_name != full_name:
                person.full_name = full_name
                changed = True
            mapped_status = _employment_status(status_raw, mapping)
            if mapped_status != EmploymentStatus.UNKNOWN and person.employment_status != mapped_status:
                person.employment_status = mapped_status
                changed = True
            if changed:
                person.save()
                audit_event(
                    actor=actor,
                    action="person.updated_from_official_source",
                    entity=person,
                    previous=previous,
                    new={"full_name": person.full_name, "employment_status": person.employment_status},
                    context={"payroll_import_id": import_record.pk, "row": row_number},
                )
            elif created:
                audit_event(
                    actor=actor,
                    action="person.created_from_official_source",
                    entity=person,
                    new={"curp": person.curp},
                    context={"payroll_import_id": import_record.pk, "row": row_number},
                )
            for system in AdministrativeSystem.objects.filter(code__in=systems, active=True):
                PersonSystemMembership.objects.get_or_create(
                    person=person,
                    system=system,
                    defaults={"first_seen_in": import_record},
                )

        row_status = SourceRowStatus.ERROR if errors else (SourceRowStatus.WARNING if warnings else SourceRowStatus.ACCEPTED)
        PayrollSourceRow.objects.create(
            payroll_import=import_record,
            row_number=row_number,
            status=row_status,
            curp_raw=raw_curp,
            normalized_curp=curp if len(curp) <= 18 else "",
            full_name=full_name,
            employment_status_raw=status_raw,
            systems_payload=systems,
            raw_payload=payload,
            errors=errors,
            warnings=warnings,
            row_fingerprint=fingerprint,
            person=person,
        )
        if errors:
            counts["errors"] += 1
        else:
            counts["accepted"] += 1
            if warnings:
                counts["warnings"] += 1

    workbook.close()
    import_record.row_count = counts["rows"]
    import_record.accepted_rows = counts["accepted"]
    import_record.error_rows = counts["errors"]
    import_record.warning_rows = counts["warnings"]
    import_record.status = (
        ImportStatus.COMPLETED_WITH_ERRORS if counts["errors"] else ImportStatus.COMPLETED
    )
    import_record.save(
        update_fields=(
            "row_count",
            "accepted_rows",
            "error_rows",
            "warning_rows",
            "status",
            "updated_at",
        )
    )
    audit_event(
        actor=actor,
        action="payroll.imported",
        entity=import_record,
        new={"sha256": sha256, **counts},
    )
    return import_record
